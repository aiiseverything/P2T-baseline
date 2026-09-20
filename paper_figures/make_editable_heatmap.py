#!/usr/bin/env python3
"""Build a five-slide, native-object PowerPoint from the saved four-case JSON.

All labels and response lines are editable text. Each token highlight is a
separate, named rectangle. No screenshot or raster image is used in the PPTX.
"""
import argparse
import contextlib
import hashlib
import io
import json
import os
from pathlib import Path
import re
import zipfile

HERE = Path(__file__).resolve().parent
os.environ.setdefault('MPLCONFIGDIR', str(HERE / '.mplconfig'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.colors import to_rgb
from matplotlib.figure import Figure
from matplotlib.font_manager import FontProperties, fontManager
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
import render_credit_case_study_original as original

STEM = 'fig6_credit_case_heatmaps_editable'
FONTS = {'Lato': ('Arial', 'Liberation Sans'),
         'Liberation Serif': ('Times New Roman', 'Liberation Serif'),
         'Liberation Mono': ('Courier New', 'Liberation Mono')}
INK, MUTED = '#29231f', '#72665e'


def rgb(color):
    return RGBColor(*(round(255*x) for x in to_rgb(color)))


def plain(text):
    return text.replace(r'$A\,w_t$', 'A·w_t').replace('$w_t = 1$', 'w_t = 1').replace('$w_t$', 'w_t')


def geometry(data):
    for path in (HERE/'fonts').glob('*.ttf'):
        fontManager.addfont(str(path))
    close, savefig = plt.close, Figure.savefig
    plt.close = lambda *a, **k: None
    Figure.savefig = lambda *a, **k: None
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            layout = original.render(data['cases'], data['rendering']['weight_max'], HERE,
                                     data['rendering']['gamma'], 144)
        fig = plt.gcf()
    finally:
        plt.close, Figure.savefig = close, savefig
    fig.canvas.draw()
    ax = fig.axes[0]
    renderer = fig.canvas.get_renderer()
    width, height = fig.get_size_inches()*72
    assert len(layout['token_boxes']) == len(data['rendering']['token_boxes']) == 76
    for actual, expected in zip(layout['token_boxes'], data['rendering']['token_boxes']):
        assert actual['case']==expected['case'] and actual['position']==expected['position']
        assert all(abs(actual[k]-expected[k])<1e-6 for k in ('x','y','width','height'))
    titles = {}
    for text in ax.texts:
        match = re.match(r'\(([a-d])\)  ', text.get_text())
        if match:
            titles[data['cases'][ord(match[1])-97]['key']] = text.get_position()
    elements=[]
    for patch, token in zip(ax.patches, layout['token_boxes']):
        case = next(c for c in data['cases'] if c['key']==token['case'])
        saved = next(t for t in case['tokens'] if t['position']==token['position'])
        elements.append(dict(kind='rect',case=case['key'],position=saved['position'],
            name=f"{case['key']} | token {saved['position']} | id {saved['token_id']} | w={saved['weight']:.6g}",
            x=patch.get_x(), y=patch.get_y(), w=patch.get_width(), h=patch.get_height(),
            color=patch.get_facecolor()[:3], edge=patch.get_edgecolor()[:3], lw=patch.get_linewidth()))
    for i, line in enumerate(ax.lines):
        x,y=line.get_data()
        assert len(x)==len(y)==2
        elements.append(dict(kind='line',name=f'Rule or legend tick {i}',x=float(x[0]),y=float(y[0]),
            x2=float(x[1]),y2=float(y[1]),color=line.get_color(),lw=line.get_linewidth()))
    for i, text in enumerate(ax.texts):
        x,y=text.get_position();content=plain(text.get_text());size=text.get_fontsize()
        original_font=text.get_fontfamily()[0];font,metric_font=FONTS[original_font]
        weight=text.get_fontweight()
        w=renderer.get_text_width_height_descent(content.replace('w_t','wt'),FontProperties(family=metric_font,size=size,weight=weight),ismath=False)[0]*72/fig.dpi
        if text.get_ha()=='center':x-=w/2
        elif text.get_ha()=='right':x-=w
        case=None
        if 90<=y<height-30:
            column=[(key,pos) for key,pos in titles.items() if (pos[0]<width/2)==(text.get_position()[0]<width/2) and pos[1]<=y+0.01]
            if column:case=max(column,key=lambda item:item[1][1])[0]
        elements.append(dict(kind='text',case=case,name=(case or 'Figure')+' | '+content[:65],
            text=content,x=float(x),y=float(y)-size*.94,w=float(w)+2,h=size*1.45,
            size=size,font=font,bold=weight in ('bold','semibold',700),color=text.get_color()))
    for image in ax.images:
        x0,x1,y1,y0=image.get_extent()
        for i in range(128):
            elements.append(dict(kind='rect',name=f'Legend color {i+1:03d}',
                x=x0+(x1-x0)*i/128,y=y0,w=(x1-x0)/128+.015,h=y1-y0,
                color=image.cmap((i+.5)/128)[:3],edge=None,lw=0))
    plt.close(fig)
    return float(width),float(height),elements,layout


def add_text(slide, element, scale, dx, dy):
    box=slide.shapes.add_textbox(Pt((element['x']*scale+dx)),Pt(element['y']*scale+dy),
        Pt(element['w']*scale),Pt(element['h']*scale))
    box.name=element['name']
    frame=box.text_frame
    frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
    frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.NONE;frame.vertical_anchor=MSO_ANCHOR.TOP
    p=frame.paragraphs[0];p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.0
    for part in re.split('(w_t)',element['text']):
        for content,subscript in ([('w',False),('t',True)] if part=='w_t' else [(part,False)]):
            if not content:continue
            run=p.add_run();run.text=content;run.font.name=element['font']
            run.font.size=Pt(element['size']*scale*(.75 if subscript else 1))
            run.font.bold=element['bold'];run.font.color.rgb=rgb(element['color'])
            if subscript:run._r.get_or_add_rPr().set('baseline','-25000')
    return box


def draw(slide,elements,scale=1,dx=0,dy=0,group_cases=True):
    cases={}
    # Rectangles must precede text so every line stays readable and continuous.
    for e in sorted(elements,key=lambda e:e['kind']=='text'):
        if e['kind']=='text':shape=add_text(slide,e,scale,dx,dy)
        elif e['kind']=='line':
            shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(e['x']*scale+dx),Pt(e['y']*scale+dy),
                Pt(e['x2']*scale+dx),Pt(e['y2']*scale+dy))
            shape.line.color.rgb=rgb(e['color']);shape.line.width=Pt(e['lw']*scale)
        else:
            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(e['x']*scale+dx),Pt(e['y']*scale+dy),
                Pt(e['w']*scale),Pt(e['h']*scale))
            shape.fill.solid();shape.fill.fore_color.rgb=rgb(e['color'])
            if e['edge'] is None:shape.line.fill.background()
            else:shape.line.color.rgb=rgb(e['edge']);shape.line.width=Pt(e['lw']*scale)
        shape.name=e['name']
        # Override the template's shape effects; token fills must stay flat.
        style=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
        if style is not None:shape._element.remove(style)
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        if e.get('case'):cases.setdefault(e['case'],[]).append(shape)
    if group_cases:
        for name,shapes in cases.items():
            slide.shapes.add_group_shape(shapes).name='CASE | '+name+' | editable text and token highlights'


def label(text,x,y,w,size=10,color=INK,bold=False):
    return dict(kind='text',name=text[:65],text=text,x=x,y=y,w=w,h=size*1.5,
                size=size,font='Arial',bold=bold,color=color)


def notes(cases,data):
    return ('EDITABLE NATIVE OBJECTS. Text lines and colored token rectangles are separate. '
        'Double-click a case group, or use Selection Pane / Ungroup to edit. '
        'If changing wording or font size, realign the highlight rectangles. '
        'Colors represent saved positive credit weights, not token reward labels. '
        'These are manually selected qualitative examples.\n\n'+
        json.dumps(dict(color_scale=data['rendering']['color_mapping'],cases=cases),ensure_ascii=False,indent=2))


def build(data_path,out_path):
    data=json.loads(data_path.read_text());width,height,elements,layout=geometry(data)
    prs=Presentation();scale=960/width
    prs.slide_width=Pt(width*scale);prs.slide_height=Pt(height*scale)
    prs.core_properties.title='VPO token-level credit allocation — editable case studies'
    prs.core_properties.subject='Native PowerPoint text and token rectangles, four saved examples'
    prs.core_properties.author='VPO-RM'
    blank=prs.slide_layouts[6]
    slide=prs.slides.add_slide(blank);draw(slide,elements,scale)
    slide.notes_slide.notes_text_frame.text=notes(data['cases'],data)
    for index,case in enumerate(data['cases']):
        slide=prs.slides.add_slide(blank)
        subset=[e for e in elements if e.get('case')==case['key']]
        x0=min(e['x'] for e in subset);y0=min(e['y'] for e in subset)
        x1=max(e['x']+e['w'] for e in subset);y1=max(e['y']+e['h'] for e in subset)
        # Reserve a separate bottom band for the legend and footer.
        zoom=min(2.25,(width-48)/(x1-x0),(height-181)/(y1-y0))
        dx=(width-(x1-x0)*zoom)/2-x0*zoom
        dy=63-y0*zoom
        draw(slide,subset,scale*zoom,scale*dx,scale*dy)
        heading=f"CASE ({chr(97+index)})  /  {case['kind'].upper()} ADVANTAGE"
        draw(slide,[label(heading,24,18,width-48,10,color='#9a4c30',bold=True),
            label('VPO · Qwen3-14B-Base · λ = 4',24,35,width-48,8,color=MUTED),
            label('Same saved token weights and shared square-root color scale as the overview.',24,height-30,width-48,7.6,color=MUTED)],scale,group_cases=False)
        # Reuse the original editable colorbar and its ticks at the lower left.
        legend=[e for e in elements if (e['name'].startswith('Legend color') or
            (e['kind']=='line' and e['x']>=351 and e['y']<45) or
            (e['kind']=='text' and e['y']<53 and e['x']>330))]
        draw(slide,legend,scale,scale*(24-352),scale*(height-92-14),group_cases=False)
        slide.notes_slide.notes_text_frame.text=notes([case],data)
    out_path.parent.mkdir(parents=True,exist_ok=True);prs.save(out_path)
    check=Presentation(out_path)
    with zipfile.ZipFile(out_path) as archive:
        assert archive.testzip() is None
        assert not any(n.startswith('ppt/media/') for n in archive.namelist())
        slide_xml=[archive.read(f'ppt/slides/slide{i}.xml') for i in range(1,6)]
        assert all(b'<p:pic>' not in xml for xml in slide_xml)
    assert len(check.slides)==5
    validation=dict(status='passed',slides=5,overview_cases=4,saved_tokens=sum(len(c['tokens']) for c in data['cases']),
        overview_token_rectangles=len(layout['token_boxes']),embedded_images=0,
        editable=['response text','titles','prompt summaries','metadata','individual token highlights','colorbar','rules'],
        fonts=['Arial','Times New Roman','Courier New'],
        source_json_sha256=hashlib.sha256(data_path.read_bytes()).hexdigest(),
        pptx_sha256=hashlib.sha256(out_path.read_bytes()).hexdigest(),
        note='Slide1 is the complete paper figure; slides2–5 are enlarged cases. Full source cases and token weights are in slide notes and the companion JSON/CSV.')
    out_path.with_suffix('.validation.json').write_text(json.dumps(validation,indent=2)+'\n')
    print(json.dumps(validation,indent=2));print(out_path)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data',type=Path,default=HERE/'fig6_credit_case_heatmaps_paper.json')
    parser.add_argument('--out',type=Path,default=HERE/(STEM+'.pptx'))
    args=parser.parse_args();build(args.data.resolve(),args.out.resolve())
